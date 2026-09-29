import sys
import os

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    C_BLUE = RGBColor(2, 132, 199)
    C_DARK = RGBColor(15, 23, 42)
    C_GRAY = RGBColor(100, 116, 139)

    def add_footer(slide):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        tf = tx_box.text_frame
        p = tf.paragraphs[0]
        p.text = "Glosario Interactivo IA • Cesar hessiel arevalo sanchez | utp0153205@alumno.utpuebla.edu.mx"
        p.font.size = Pt(10)
        p.font.color.rgb = C_GRAY

    # Slide 1: Portada
    slide1 = prs.slides.add_slide(blank_layout)
    tx1 = slide1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(3.5))
    tf1 = tx1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "EVOLUCIÓN, GLOSARIO Y CLASIFICACIÓN DE LA IA"
    p1.font.bold = True
    p1.font.size = Pt(36)
    p1.font.color.rgb = C_DARK

    p2 = tf1.add_paragraph()
    p2.text = "Presentación Ejecutiva e Infografía Interactiva"
    p2.font.size = Pt(20)
    p2.font.color.rgb = C_BLUE
    p2.space_before = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "Autor: Cesar hessiel arevalo sanchez\nContacto: utp0153205@alumno.utpuebla.edu.mx"
    p3.font.size = Pt(14)
    p3.font.color.rgb = C_GRAY
    p3.space_before = Pt(30)
    add_footer(slide1)

    # Slide 2: Clasificación de IA
    slide2 = prs.slides.add_slide(blank_layout)
    tx2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(1.0))
    p = tx2.text_frame.paragraphs[0]
    p.text = "CLASIFICACIÓN CLÁSICA DE LA INTELIGENCIA ARTIFICIAL"
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = C_DARK

    tipos = [
        ("IA DÉBIL (Narrow AI)", "Dominante hoy", "Especializada en una sola tarea sin conciencia general. Ejemplos: Recomendadores, Traductores, Chatbots acotados."),
        ("IA FUERTE (AGI)", "En Desarrollo", "Capacidad cognitiva general equivalente a la humana. Adaptabilidad entre dominios sin reentrenamiento."),
        ("IA SUPERINTELIGENTE (ASI)", "Frontera Teórica", "Supera con creces la capacidad intelectual humana combinada en creatividad, ciencia y resolución de problemas.")
    ]

    for i, (t, badge, d) in enumerate(tipos):
        left = Inches(0.8 + i * 3.9)
        shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.7), Inches(4.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = C_BLUE
        
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = t
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = C_DARK

        p_sub = tf.add_paragraph()
        p_sub.text = f"[{badge}]"
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = C_BLUE
        p_sub.space_before = Pt(5)

        p_desc = tf.add_paragraph()
        p_desc.text = d
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = C_GRAY
        p_desc.space_before = Pt(15)

    add_footer(slide2)

    # Slide 3: Línea de Tiempo
    slide3 = prs.slides.add_slide(blank_layout)
    tx3 = slide3.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(1.0))
    p = tx3.text_frame.paragraphs[0]
    p.text = "LÍNEA DE TIEMPO: EVOLUCIÓN HISTÓRICA (10 HITOS)"
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = C_DARK

    hitos_brief = [
        ("1950s", "Pioneros (Turing / Dartmouth)"),
        ("1960s", "ELIZA chatbot"),
        ("1970s", "Primer Invierno de IA"),
        ("1980s", "Sistemas Expertos"),
        ("1990s", "Deep Blue vs Kasparov"),
        ("2000s", "Auge de Machine Learning"),
        ("2010s", "Deep Learning & AlphaGo"),
        ("2020s", "GPT & ChatGPT"),
        ("2023-24", "IA Multimodal"),
        ("Hoy", "Agentes Autónomos")
    ]

    for i, (periodo, dec) in enumerate(hitos_brief):
        row = i // 5
        col = i % 5
        left = Inches(0.8 + col * 2.35)
        top = Inches(1.8 + row * 2.4)
        
        shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.2), Inches(2.1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(241, 245, 249)
        shape.line.color.rgb = C_BLUE

        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = periodo
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = C_BLUE

        p2 = tf.add_paragraph()
        p2.text = dec
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK
        p2.space_before = Pt(8)

    add_footer(slide3)

    static_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")
    os.makedirs(static_dir, exist_ok=True)
    out_path = os.path.join(static_dir, "Glosario_y_Evolucion_IA.pptx")
    prs.save(out_path)
    print(f"Presentación PPTX creada exitosamente en: {out_path}")

except Exception as e:
    print(f"Aviso al crear PPTX: {e}")
